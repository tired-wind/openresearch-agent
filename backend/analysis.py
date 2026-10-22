"""Predefined analysis tools. Nothing from a question is evaluated as code."""
from __future__ import annotations
import csv
import io
import re
import zipfile
import hashlib
import numpy as np
import pandas as pd
from sklearn.compose import ColumnTransformer
from sklearn.dummy import DummyClassifier, DummyRegressor
from sklearn.ensemble import RandomForestClassifier, RandomForestRegressor
from sklearn.impute import SimpleImputer
from sklearn.inspection import permutation_importance
from sklearn.metrics import balanced_accuracy_score, f1_score, roc_auc_score, average_precision_score, mean_absolute_error, root_mean_squared_error, r2_score
from sklearn.model_selection import train_test_split, cross_val_score, StratifiedKFold, KFold
from sklearn.pipeline import Pipeline
from sklearn.preprocessing import OneHotEncoder, LabelEncoder

SEED = 42
MAX_BYTES = 5 * 1024 * 1024

class InputError(ValueError):
    pass

def read_table(raw: bytes, suffix: str) -> pd.DataFrame:
    if not raw or len(raw) > MAX_BYTES:
        raise InputError('File must be nonempty and at most 5 MiB.')
    try:
        if suffix == '.csv':
            text = raw.decode('utf-8-sig')
            reader = csv.reader(io.StringIO(text))
            header = next(reader)
            for row_index, row in enumerate(reader):
                if row and len(row) != len(header):
                    raise InputError('CSV row has inconsistent column count.')
                if row_index >= 5000:
                    raise InputError('CSV exceeds 5,000 rows.')
            df = pd.read_csv(io.StringIO(text), nrows=5001)
        elif suffix == '.xlsx':
            with zipfile.ZipFile(io.BytesIO(raw)) as z:
                if sum(i.file_size for i in z.infolist()) > 30 * 1024 * 1024 or len(z.infolist()) > 500:
                    raise InputError('Workbook expands beyond the safe size limit.')
                if any('externalLinks/' in i.filename for i in z.infolist()):
                    raise InputError('Workbook external links are not supported.')
            import openpyxl
            book = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
            try:
                sheet = book.worksheets[0]
                if (sheet.max_row or 0) > 5001 or (sheet.max_column or 0) > 60:
                    raise InputError('Workbook exceeds 5,000 rows or 60 columns.')
                rows = []
                for row in sheet.iter_rows():
                    if any(cell.data_type == 'f' for cell in row):
                        raise InputError('Formula cells are unsupported; upload a values-only workbook.')
                    rows.append([cell.value for cell in row])
                    if len(rows) > 5001 or len(row) > 60:
                        raise InputError('Workbook exceeds row or column limits.')
                if not rows:
                    raise InputError('Workbook is empty.')
                header, *body = rows
                df = pd.DataFrame(body, columns=header)
            finally:
                book.close()
        else:
            raise InputError('Only UTF-8 CSV and values-only XLSX are supported.')
        names = [str(x).strip() if x is not None else '' for x in header]
        if any(not x or len(x) > 100 or any(ord(c) < 32 for c in x) for x in names) or len(set(names)) != len(names):
            raise InputError('Column names must be nonempty, unique and at most 100 characters.')
        if len(names) != len(df.columns):
            raise InputError('Inconsistent column count.')
        df.columns = names
        if not 2 <= len(df.columns) <= 60 or not 1 <= len(df) <= 5000:
            raise InputError('Supported size: 1–5,000 rows and 2–60 columns.')
        if df.select_dtypes(include='object').map(lambda x: len(str(x)) > 2000).any().any():
            raise InputError('Cell text exceeds 2,000 characters.')
        for col in df.select_dtypes(include='object'):
            df[col] = df[col].map(lambda v: str(v) if pd.notna(v) else np.nan)
        return df.replace([np.inf, -np.inf], np.nan)
    except InputError:
        raise
    except Exception as exc:
        raise InputError('Could not parse file; check encoding, table shape and workbook format.') from exc

def overview(df: pd.DataFrame) -> dict:
    return {'rows': len(df), 'columns': len(df.columns), 'duplicate_rows': int(df.duplicated().sum()),
            'missing_cells': int(df.isna().sum().sum()),
            'fields': [{'name': str(c), 'dtype': str(df[c].dtype), 'missing': int(df[c].isna().sum()), 'unique': int(df[c].nunique())} for c in df],
            'numeric_summary': df.describe().replace([np.inf, -np.inf], np.nan).round(5).to_dict() if len(df.select_dtypes('number').columns) else {}}

def analyze(df: pd.DataFrame, target: str, task: str = 'auto') -> dict:
    if target not in df:
        raise InputError('Target column does not exist.')
    original_n = len(df)
    duplicates = int(df.duplicated().sum())
    missing_target = int(df[target].isna().sum())
    df = df.drop_duplicates().dropna(subset=[target]).reset_index(drop=True)
    if len(df) < 60:
        raise InputError('At least 60 complete-target, distinct rows are required.')
    y = df[target]
    if task == 'auto':
        task = 'classification' if not pd.api.types.is_numeric_dtype(y) or y.nunique() <= 10 else 'regression'
    if task == 'classification':
        y = y.astype(str)
        if not 2 <= y.nunique() <= 10 or y.value_counts().min() < 8:
            raise InputError('Classification requires 2–10 classes with at least 8 rows each.')
    elif not pd.api.types.is_numeric_dtype(y) or y.nunique() < 10:
        raise InputError('Regression requires a numeric target with at least 10 distinct values.')
    X = df.drop(columns=[target])
    train, test = train_test_split(np.arange(len(df)), test_size=.25, random_state=SEED, stratify=y if task == 'classification' else None)
    Xtr, Xte, ytr, yte = X.iloc[train].copy(), X.iloc[test].copy(), y.iloc[train], y.iloc[test]
    excluded = []
    for c in Xtr:
        s = Xtr[c]
        reason = None
        if re.search(r'(^|[_\s])(id|uuid|name|email|phone|address|date|timestamp)([_\s]|$)', c, re.I):
            reason = 'Identifier, direct personal field or time field; review subject grouping and temporal design.'
        elif s.nunique() <= 1:
            reason = 'Constant or entirely missing in training rows.'
        elif not pd.api.types.is_numeric_dtype(s) and s.nunique() > min(40, len(s)*.4):
            reason = 'High-cardinality text/categorical feature.'
        elif s.notna().all() and s.astype(str).equals(ytr.astype(str)):
            reason = 'Exact target copy in training rows.'
        elif task == 'regression' and pd.api.types.is_numeric_dtype(s) and abs(s.corr(ytr)) > .9999:
            reason = 'Near-perfect target proxy in training rows; possible leakage.'
        elif task == 'classification' and s.nunique() == ytr.nunique() and s.notna().all():
            pairs = pd.DataFrame({'x':s.astype(str), 'y':ytr})
            if pairs.groupby('x').y.nunique().max() == 1:
                reason = 'One-to-one target encoding in training rows; possible leakage.'
        if reason:
            excluded.append({'feature':c,'reason':reason})
    keep = [c for c in Xtr if c not in {x['feature'] for x in excluded}]
    if not keep:
        raise InputError('No usable features remain after identifier, constant and leakage checks.')
    Xtr, Xte = Xtr[keep], Xte[keep]
    nums = list(Xtr.select_dtypes('number').columns)
    cats = [c for c in keep if c not in nums]
    transforms = []
    if nums:
        transforms.append(('numeric', SimpleImputer(strategy='median', keep_empty_features=True), nums))
    if cats:
        transforms.append(('categorical', Pipeline([('impute', SimpleImputer(strategy='most_frequent', keep_empty_features=True)), ('encode', OneHotEncoder(handle_unknown='ignore', sparse_output=False))]), cats))
    prep = ColumnTransformer(transforms)
    classifier = task == 'classification'
    label_encoder = None
    if classifier:
        label_encoder = LabelEncoder().fit(ytr)
        ytr = pd.Series(label_encoder.transform(ytr), index=ytr.index)
        yte = pd.Series(label_encoder.transform(yte), index=yte.index)
    estimator = RandomForestClassifier(n_estimators=100, max_depth=6, min_samples_leaf=3, class_weight='balanced', random_state=SEED, n_jobs=1) if classifier else RandomForestRegressor(n_estimators=100, max_depth=6, min_samples_leaf=3, random_state=SEED, n_jobs=1)
    pipe = Pipeline([('preprocess', prep), ('model', estimator)])
    scoring = 'balanced_accuracy' if classifier else 'neg_root_mean_squared_error'
    folds = StratifiedKFold(3, shuffle=True, random_state=SEED) if classifier else KFold(3, shuffle=True, random_state=SEED)
    cv = cross_val_score(pipe, Xtr, ytr, cv=folds, scoring=scoring, error_score='raise')
    pipe.fit(Xtr, ytr)
    pred = pipe.predict(Xte)
    if classifier:
        dummy = DummyClassifier(strategy='prior').fit(np.zeros((len(ytr),1)), ytr)
        metrics = {'balanced_accuracy': float(balanced_accuracy_score(yte,pred)), 'f1_macro':float(f1_score(yte,pred,average='macro',zero_division=0)), 'baseline_balanced_accuracy':float(balanced_accuracy_score(yte,dummy.predict(np.zeros((len(yte),1)))))}
        if y.nunique() == 2:
            prob = pipe.predict_proba(Xte)[:,1]
            positive_code = pipe.classes_[1]
            positive = str(label_encoder.inverse_transform([positive_code])[0])
            metrics.update(roc_auc=float(roc_auc_score(yte==positive_code,prob)), average_precision=float(average_precision_score(yte==positive_code,prob)))
        else:
            positive = None
    else:
        dummy = DummyRegressor().fit(np.zeros((len(ytr),1)), ytr)
        metrics = {'rmse':float(root_mean_squared_error(yte,pred)), 'mae':float(mean_absolute_error(yte,pred)), 'r2':float(r2_score(yte,pred)), 'baseline_rmse':float(root_mean_squared_error(yte,dummy.predict(np.zeros((len(yte),1)))))}
        positive = None
    pi = permutation_importance(pipe, Xte, yte, scoring=scoring, n_repeats=8, random_state=SEED, n_jobs=1)
    factors = sorted([{'feature':c, 'importance':float(pi.importances_mean[i]), 'repeat_std':float(pi.importances_std[i]), 'evidence_id':f'factor-{i+1}', 'evidence_type':'computation'} for i,c in enumerate(keep)],key=lambda r:r['importance'],reverse=True)[:10]
    limitations = ['Predictive importance is not causal evidence and does not establish an association direction.', 'Permutation repeat SD measures shuffle variability, not a confidence interval.', 'Single random holdout; small-sample performance and rankings can be unstable.', 'No temporal or grouped validation. Independent rows must be confirmed by the user.', 'Leakage heuristics cannot detect all post-outcome variables; domain review is required.', 'Feature rankings use the evaluation set and are exploratory; new validation is required before selecting/refitting features.', 'No clinical validation, treatment advice, p-values or invented confidence intervals.']
    if duplicates or missing_target:
        limitations.append(f'Dropped {duplicates} exact duplicates before splitting; original missing-target rows: {missing_target}.')
    return {'task':task, 'target':target, 'n_original':original_n, 'n_analyzed':len(df), 'n_train':len(train), 'n_test':len(test), 'positive_class':positive,
            'seed':SEED,'model':'RandomForest (100 trees, max_depth=6, min_samples_leaf=3)', 'metrics':metrics,
            'cv':{'folds':3,'scope':'training rows only; preprocessing refitted inside every fold','scoring':scoring, 'scores':[float(x) for x in cv], 'mean':float(cv.mean()),'std':float(cv.std())},
            'factors':factors,'excluded':excluded,'features_used':keep,'importance_scoring':scoring,'limitations':limitations,
            'split':{'train_indices':train.tolist(),'test_indices':test.tolist(),'index_scope':'after exact deduplication and missing-target removal'},
            'evidence_type':'computation', 'evidence_id':'analysis-1'}
